"""Contrôles indépendants de l'auditeur (contre-audit de main d745ead, 2026-10-23).

Écrits SANS reprendre les tests existants du dépôt : chaque contrôle parle à
l'API HTTP réelle et relit la base par SQL. Ils couvrent D2, D4, D6, D7, D9,
C12-02 (phase 1) puis C7 et la concurrence (phase 2).

Préparation (base JETABLE, jamais quincaillerie_test ni une base réelle) :

    createdb qf_chk
    psql -d qf_chk -f QuincaillerieFranck_Test/creation_base_donnees.sql
    PGDATABASE=qf_chk PSQL=psql bash db/outils/migrer.sh appliquer
    psql -d qf_chk -f db/tests/00_jeu_essai.sql
    psql -d qf_chk -v mdp="'qf_app_dev_local'" -f db/outils/definir_mot_de_passe_app.sql
    psql -d qf_chk -c "UPDATE utilisateurs SET tentatives_echouees = 0,
        mot_de_passe_hash = crypt('Pw-' || identifiant || '-1', gen_salt('bf', 12))"
    # config.ini de TEST (dbname = qf_chk) puis :
    QF_CONFIG=<config_chk.ini> QF_FICHIER_ETAT_SAUVEGARDE=<dossier>/etat.json \\
        python -m uvicorn app.main:app --app-dir server --port 8020 &
    QF_AUDIT_ETAT=<dossier>/etat.json python controles_independants.py

Variables : QF_AUDIT_URL (défaut http://127.0.0.1:8020), QF_AUDIT_DB (qf_chk),
QF_AUDIT_ETAT (fichier d'état de sauvegarde lu par le serveur), PGHOST/PGPORT/
PGUSER/PGPASSWORD pour psql. Sortie : une ligne OK/FAIL par contrôle ; code de
sortie 1 s'il y a au moins un FAIL. Un FAIL n'est pas forcément une régression :
plusieurs contrôles consignent volontairement un DÉFAUT NOUVEAU (voir
AUDIT_INDEPENDANT_d745ead.md) et resteront FAIL tant que ce défaut vit.
"""

from __future__ import annotations

import base64
import concurrent.futures as cf
import datetime
import hashlib
import hmac
import io
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request

BASE = os.environ.get("QF_AUDIT_URL", "http://127.0.0.1:8020")
DB = os.environ.get("QF_AUDIT_DB", "qf_chk")
ETAT = os.environ.get("QF_AUDIT_ETAT", "")
CONFIG = os.environ.get("QF_CONFIG", "")
os.environ.setdefault("PGHOST", "127.0.0.1")
os.environ.setdefault("PGPORT", "5433")
os.environ.setdefault("PGUSER", "postgres")
os.environ.setdefault("PGPASSWORD", "qf_dev_local")

RESULTATS: list[tuple[str, bool, str]] = []
HORLOGE = int(time.time())
COMPTEUR = [0]


def sql(requete: str) -> str:
    r = subprocess.run(
        ["psql", "-d", DB, "-Atq", "-F", "|", "-c", requete], capture_output=True, text=True
    )
    return (r.stdout + r.stderr).strip()


def api(methode, chemin, corps=None, jeton=None, brut=False):
    req = urllib.request.Request(
        BASE + chemin,
        method=methode,
        data=json.dumps(corps).encode() if corps is not None else None,
        headers={"Content-Type": "application/json", **({"Authorization": "Bearer " + jeton} if jeton else {})},
    )
    try:
        with urllib.request.urlopen(req) as r:
            donnees = r.read()
            if brut:
                return r.status, donnees
            return r.status, (json.loads(donnees) if donnees[:1] in (b"{", b"[") else donnees)
    except urllib.error.HTTPError as e:
        donnees = e.read()
        try:
            return e.code, json.loads(donnees)
        except Exception:  # noqa: BLE001
            return e.code, donnees


def connexion(identifiant: str) -> str:
    code, corps = api("POST", "/auth/connexion", {"identifiant": identifiant, "mot_de_passe": f"Pw-{identifiant}-1"})
    assert code == 200, (identifiant, code, corps)
    return corps["jeton"]


def verif(nom: str, ok: bool, detail="") -> None:
    RESULTATS.append((nom, bool(ok), str(detail)))
    print(("OK   " if ok else "FAIL ") + nom + (f" :: {detail}" if detail and not ok else ""))


def facturier(prefixe="MAG-AUD") -> str:
    COMPTEUR[0] += 1
    return f"{prefixe}{HORLOGE}{COMPTEUR[0]}"


def vente(jeton, corps, prefixe="MAG-AUD"):
    base = {"numero_facturier": facturier(prefixe), "vendeur_id": 1, **corps}
    return api("POST", "/ventes", base, jeton)


def ligne(article=1, **k):
    return {"article_id": article, "quantite": 1, **k}


def id_vente(corps):
    return corps.get("vente_id") if isinstance(corps, dict) else None


def main() -> int:
    sql("TRUNCATE ventes CASCADE")
    sql("TRUNCATE comptages_stock CASCADE")
    sql("UPDATE stocks_sites SET quantite_stock = 1000 WHERE site_id = 1")
    resp, compta, comptoir = connexion("resp"), connexion("magasin.compta"), connexion("comptoir.compta")
    resp1 = lambda **k: vente(resp, {"site_id": 1, **k})  # noqa: E731

    # ------------------------------------------------------------------ D2
    c, b = resp1(lignes=[ligne(prix_unitaire=6500)], paiements=[
        {"mode_paiement": "especes", "montant": 4000}, {"mode_paiement": "orange_money", "montant": 2500}])
    verif("D2 mixte 4000+2500 accepté", c == 201, (c, b))
    v = id_vente(b)
    if v:
        verif("D2 mode stocké = mixte", sql(f"SELECT mode_paiement FROM ventes WHERE id={v}") == "mixte")
        verif("D2 composantes tracées par mode", sql(
            f"SELECT string_agg(mode_paiement||':'||montant::int, ',' ORDER BY mode_paiement) "
            f"FROM ventes_paiements WHERE vente_id={v}") == "especes:4000,orange_money:2500")
        verif("D2 une seule recette du total", sql(
            f"SELECT count(*)||':'||sum(montant)::int FROM transactions WHERE vente_id={v} AND type='recette'") == "1:6500")
    c, _ = resp1(lignes=[ligne(prix_unitaire=6500)], paiements=[
        {"mode_paiement": "especes", "montant": 4000}, {"mode_paiement": "orange_money", "montant": 2000}])
    verif("D2 somme ≠ total refusée (422)", c == 422, c)
    c, _ = resp1(lignes=[ligne(prix_unitaire=6500)], mode_paiement="especes",
                 paiements=[{"mode_paiement": "especes", "montant": 6500}])
    verif("D2 mode_paiement + paiements simultanés refusés", c == 422, c)
    c, _ = resp1(lignes=[ligne(prix_unitaire=6500)], paiements=[
        {"mode_paiement": "especes", "montant": 3000}, {"mode_paiement": "especes", "montant": 3500}])
    verif("D2 même mode deux fois : refus propre, pas de 500", c in (400, 409, 422), c)
    c, _ = resp1(lignes=[ligne(prix_unitaire=6500)], paiements=[{"mode_paiement": "credit_client", "montant": 6500}])
    verif("D2 composante crédit refusée", c == 422, c)
    c, _ = resp1(lignes=[ligne(prix_unitaire=6500)], paiements=[
        {"mode_paiement": "especes", "montant": -5}, {"mode_paiement": "orange_money", "montant": 6505}])
    verif("D2 composante négative refusée", c == 422, c)
    sortie = sql("""BEGIN; WITH v AS (INSERT INTO ventes (site_id,utilisateur_id,statut,mode_paiement,numero_facturier,
        vendeur_id,sous_total_ht,montant_tva,total_ttc,utilisateur_caisse_id,date_encaissement)
        VALUES (1,4,'payee','mixte','MAG-SQLX',1,100,0,100,4,now()) RETURNING id)
        INSERT INTO ventes_paiements(vente_id,mode_paiement,montant) SELECT id,'especes',40 FROM v; COMMIT;""")
    verif("D2 base : composantes ≠ total rejetées même en SQL direct", "incohérents" in sortie, sortie[:160])
    jour = datetime.date.today().isoformat()
    c, b = api("GET", f"/caisse/attendu?site_id=1&date_cloture={jour}", jeton=resp)
    verif("D2 caisse : attendu ventilé par mode", c == 200 and b.get("attendu_especes", 0) >= 4000
          and b.get("attendu_orange_money", 0) >= 2500, b)

    # ------------------------------------------------------------------ D4
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(remise_pct=10)]})
    verif("D4 agent : −10 % ligne accepté (seuil strict >)", c == 201, c)
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(remise_pct=10.01)]})
    verif("D4 agent : −10,01 % ligne → 403", c == 403, c)
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(remise_montant=651)]})
    verif("D4 agent : remise_montant 651/6500 → 403", c == 403, c)
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(prix_unitaire=5800)]})
    verif("D4 agent : contournement par prix_unitaire 5800 → 403", c == 403, c)
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(prix_unitaire=6500)], "remise_globale_pct": 10.5})
    verif("D4 agent : remise globale 10,5 % → 403", c == 403, c)
    c, _ = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(prix_unitaire=6500)], "remise_globale_montant": 700})
    verif("D4 agent : remise globale 700/6500 → 403", c == 403, c)
    c, _ = resp1(mode_paiement="especes", lignes=[ligne(remise_pct=50)])
    verif("D4 responsable : −50 % accepté", c == 201, c)
    c, b = vente(compta, {"mode_paiement": "especes", "lignes": [ligne(remise_pct=10)], "remise_globale_pct": 10})
    verif("D4 [DÉFAUT NOUVEAU] cumul ligne −10 % + globale −10 % (≈ −19 %) par un agent doit exiger le responsable",
          c == 403, f"HTTP {c} — vente acceptée à {sql('SELECT total_ttc FROM ventes ORDER BY id DESC LIMIT 1')} FCFA sur 6500")

    # ------------------------------------------------------------------ D9
    c, b = vente(compta, {"mode_paiement": "especes", "utilisateur_caisse_id": 1, "lignes": [ligne(prix_unitaire=6500)]})
    v = id_vente(b)
    verif("D9 injection utilisateur_caisse_id ignorée : encaisseur = agent authentifié",
          c == 201 and sql(f"SELECT utilisateur_caisse_id FROM ventes WHERE id={v}") == "4", (c, b))
    c, b = resp1(mode_paiement="especes", lignes=[ligne(prix_unitaire=6500)])
    verif("D9 responsable encaisse = lui-même", c == 201 and sql(f"SELECT utilisateur_caisse_id FROM ventes WHERE id={id_vente(b)}") == "1")
    c, b = vente(comptoir, {"mode_paiement": "especes", "vendeur_id": 2,
                            "lignes": [{"article_id": 3, "quantite": 1, "prix_unitaire": 800}]}, "CPT-AUD")
    verif("D9 agent du comptoir : encaisseur = id 5", c == 201 and sql(f"SELECT utilisateur_caisse_id FROM ventes WHERE id={id_vente(b)}") == "5", (c, b))

    # ------------------------------------------------------------------ D6
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "..", "server"))
    cle = b"0" * 64
    if CONFIG:
        for l in open(CONFIG, encoding="utf-8"):
            if l.strip().startswith("secret_key"):
                cle = l.split("=", 1)[1].strip().encode()
    b64 = lambda x: base64.urlsafe_b64encode(x).rstrip(b"=")  # noqa: E731

    def forger(charge, k=None):
        corps = b64(json.dumps(charge, separators=(",", ":")).encode())
        return (corps + b"." + b64(hmac.new(k or cle, corps, hashlib.sha256).digest())).decode()

    maintenant = int(time.time())
    base_charge = {"uid": 1, "role": "responsable", "roles": ["responsable"], "site": None, "nom": "x", "chg": False}
    code, corps = api("POST", "/auth/connexion", {"identifiant": "resp", "mot_de_passe": "Pw-resp-1"})
    charge = json.loads(base64.urlsafe_b64decode(corps["jeton"].split(".")[0] + "=="))
    verif("D6 jeton réel : exp − iat = 480 min", charge["exp"] - charge["iat"] == 28800, charge["exp"] - charge["iat"])
    c, _ = api("GET", "/moi", jeton=forger({**base_charge, "iat": maintenant, "exp": maintenant + 600 * 60, "jti": "a" * 32}))
    verif("D6 jeton correctement signé mais 600 min → 401", c == 401, c)
    c, _ = api("GET", "/moi", jeton=forger({**base_charge, "iat": maintenant - 9000, "exp": maintenant - 10, "jti": "b" * 32}))
    verif("D6 jeton expiré → 401", c == 401, c)
    c, _ = api("GET", "/moi", jeton=forger({**base_charge, "iat": maintenant, "exp": maintenant + 3600, "jti": "c" * 32}))
    verif("D6 jeton signé mais session non enregistrée en base → 401", c == 401, c)
    c, _ = api("GET", "/moi", jeton=forger({**base_charge, "iat": maintenant, "exp": maintenant + 3600, "jti": "d" * 32}, b"autre-cle"))
    verif("D6 mauvaise signature → 401", c == 401, c)
    t = connexion("resp")
    verif("D6 session normale → 200", api("GET", "/moi", jeton=t)[0] == 200)
    sql("UPDATE sessions_serveur SET expire_a = now() - interval '1 second' WHERE utilisateur_id = 1")
    verif("D6 expiration côté serveur (session échue) → 401", api("GET", "/moi", jeton=t)[0] == 401)
    t = connexion("resp")
    c1, _ = api("POST", "/auth/deconnexion", jeton=t)
    verif("D6 révocation par déconnexion", c1 == 204 and api("GET", "/moi", jeton=t)[0] == 401)
    t1, t2 = connexion("magasin.stock"), connexion("magasin.stock")
    sql("UPDATE utilisateurs SET actif=false WHERE identifiant='magasin.stock'")
    verif("D6 désactivation du compte coupe toutes ses sessions",
          api("GET", "/moi", jeton=t1)[0] == 401 and api("GET", "/moi", jeton=t2)[0] == 401)
    sql("UPDATE utilisateurs SET actif=true WHERE identifiant='magasin.stock'")
    t = connexion("magasin.stock")
    sql("UPDATE utilisateurs SET mot_de_passe_hash = crypt('Pw-magasin.stock-1', gen_salt('bf',12)) WHERE identifiant='magasin.stock'")
    verif("D6 réinitialisation du mot de passe coupe la session en cours", api("GET", "/moi", jeton=t)[0] == 401)

    # ------------------------------------------------------------------ D7
    sql("DELETE FROM utilisateurs WHERE identifiant='cumul.sc'")
    sql("INSERT INTO utilisateurs(nom_complet,identifiant,mot_de_passe_hash,role,site_id,doit_changer_mot_de_passe) "
        "VALUES ('Cumul','cumul.sc',crypt('Pw-cumul.sc-1',gen_salt('bf',12)),'agent_stock',1,false)")
    sql("INSERT INTO utilisateurs_roles(utilisateur_id,role) SELECT id,r FROM utilisateurs, "
        "unnest(array['agent_stock','agent_comptabilite']) r WHERE identifiant='cumul.sc'")
    t = connexion("cumul.sc")
    c, b = api("GET", "/moi", jeton=t)
    verif("D7 /moi expose l'union des rôles", sorted(b.get("roles", [])) == ["agent_comptabilite", "agent_stock"], b)
    c, _ = api("POST", "/ventes", {"numero_facturier": facturier(), "vendeur_id": 1, "mode_paiement": "especes",
                                    "lignes": [ligne(prix_unitaire=6500)]}, t)
    verif("D7 compte cumulé : vente (droit comptabilité) accepté", c == 201, c)
    verif("D7 compte cumulé : /admin/comptes reste interdit", api("GET", "/admin/comptes", jeton=t)[0] == 403)
    sql("DELETE FROM utilisateurs_roles WHERE role='agent_comptabilite' AND utilisateur_id=(SELECT id FROM utilisateurs WHERE identifiant='cumul.sc')")
    c, _ = api("POST", "/ventes", {"numero_facturier": facturier(), "vendeur_id": 1, "mode_paiement": "especes",
                                    "lignes": [ligne(prix_unitaire=6500)]}, t)
    verif("D7 rôle retiré en base : effet immédiat sur la session en cours (403)", c == 403, c)

    # ------------------------------------------------------------------ C12-02
    if ETAT:
        def etat(contenu):
            if contenu is None:
                if os.path.exists(ETAT):
                    os.remove(ETAT)
            else:
                open(ETAT, "w", encoding="utf-8").write(contenu)
            return api("GET", "/exploitation/etat-sauvegarde", jeton=connexion("resp"))

        c, b = etat('{"resultat":"SUCCES","horodatage":"2026-10-23T03:00:00"}')
        verif("C12-02 succès daté → succes", b.get("etat") == "succes", b)
        c, b = etat('{"resultat":"SUCCES"}')
        verif("C12-02 succès SANS date → inconnu (jamais de succès sans preuve)", b.get("etat") == "inconnu", b)
        c, b = etat('{"resultat":"ECHEC","horodatage":"2026-10-23T03:00:00","messages":["disque plein"]}')
        verif("C12-02 échec → echec avec motif", b.get("etat") == "echec" and "disque plein" in b.get("motif", ""), b)
        c, b = etat(None)
        verif("C12-02 fichier absent → absent", b.get("etat") == "absent", b)
        c, b = etat("{pas du json")
        verif("C12-02 JSON corrompu → inconnu", b.get("etat") == "inconnu", b)
        c, b = etat("[]")
        verif("C12-02 [DÉFAUT NOUVEAU] JSON valide mais non-objet ([]) → pas de 500", c == 200, (c, b))
        c, b = etat('"SUCCES"')
        verif("C12-02 [DÉFAUT NOUVEAU] JSON chaîne → pas de 500", c == 200, (c, b))
        c, b = etat('{"resultat":"SUCCES","horodatage":"2024-01-01T03:00:00"}')
        verif("C12-02 [DÉFAUT NOUVEAU] dernier succès vieux de ~2 ans : alerte de péremption attendue",
              b.get("etat") != "succes", b)
        verif("C12-02 agent → 403", api("GET", "/exploitation/etat-sauvegarde", jeton=compta)[0] == 403)
        verif("C12-02 sans jeton → 401", api("GET", "/exploitation/etat-sauvegarde")[0] == 401)
        etat('{"resultat":"SUCCES","horodatage":"2026-10-23T03:00:00"}')

    # ------------------------------------------------------------------ C7
    # Le contrôle D6 a volontairement échu les sessions de `resp` : on se reconnecte.
    resp = connexion("resp")
    stock, cpt = connexion("magasin.stock"), connexion("comptoir.stock")
    c, b = api("GET", "/inventaire/articles-a-compter?moment=matin", jeton=stock)
    verif("C7 liste à compter sans quantité attendue ni écart",
          c == 200 and not any(k in json.dumps(b) for k in ("quantite_stock", "quantite_attendue", "ecart")))
    c, b = api("POST", "/inventaire/comptages", {"article_id": 1, "moment": "matin", "quantite_comptee": 25}, stock)
    verif("C7 comptage aveugle : réponse sans attendu ni écart", c == 201 and "quantite_attendue" not in b and "ecart" not in b, (c, b))
    verif("C7 doublon article/moment/jour → 409",
          api("POST", "/inventaire/comptages", {"article_id": 1, "moment": "matin", "quantite_comptee": 25}, stock)[0] == 409)
    verif("C7 au-dessus du plafond (10 000) → 422",
          api("POST", "/inventaire/comptages", {"article_id": 1, "moment": "soir", "quantite_comptee": 10001}, stock)[0] == 422)
    verif("C7 au plafond (10 000) accepté",
          api("POST", "/inventaire/comptages", {"article_id": 1, "moment": "soir", "quantite_comptee": 10000}, stock)[0] == 201)
    verif("C7 quantité négative → 422",
          api("POST", "/inventaire/comptages", {"article_id": 2, "moment": "matin", "quantite_comptee": -1}, stock)[0] == 422)
    c, b = api("POST", "/inventaire/comptages", {"article_id": 2, "moment": "matin", "quantite_comptee": 2.5}, stock)
    verif("C7 [DÉFAUT NOUVEAU] décimal sur article entier → 422 (et non 500)", c == 422, (c, b))
    verif("C7 article d'un autre site → 422",
          api("POST", "/inventaire/comptages", {"article_id": 3, "moment": "matin", "quantite_comptee": 5}, stock)[0] == 422)
    c, b = api("POST", "/inventaire/comptages", {"article_id": 3, "moment": "soir", "quantite_comptee": 5, "site_id": 1}, cpt)
    verif("C7 site_id injecté ignoré pour un agent", c == 201 and sql(f"SELECT site_id FROM comptages_stock WHERE id={b.get('comptage_id')}") == "2", (c, b))
    verif("C7 agent stock : /inventaire/ecarts interdit", api("GET", "/inventaire/ecarts", jeton=stock)[0] == 403)
    c, b = api("GET", "/inventaire/ecarts", jeton=resp)
    ecart = next((x for x in b.get("ecarts", []) if x["quantite_comptee"] == 25), None)
    verif("C7 responsable voit l'écart, égal à comptée − attendue",
          ecart is not None and float(ecart["ecart"]) == float(ecart["quantite_comptee"]) - float(ecart["quantite_attendue"]), b)
    if ecart:
        avant = sql("SELECT sum(quantite_stock) FROM stocks_sites")
        c, _ = api("POST", f"/inventaire/ecarts/{ecart['id']}/regulariser", {"type_resolution": "vol_presume", "motif": "audit"}, resp)
        verif("C7 régularisation → 201", c == 201, c)
        c, _ = api("POST", f"/inventaire/ecarts/{ecart['id']}/regulariser", {"type_resolution": "vol_presume", "motif": "bis"}, resp)
        verif("C7 seconde régularisation du même écart refusée (pas 500)", c in (409, 422), c)
        verif("C7 la régularisation ne touche jamais le stock", sql("SELECT sum(quantite_stock) FROM stocks_sites") == avant)
        verif("C7 régularisation par un agent → 403",
              api("POST", f"/inventaire/ecarts/{ecart['id']}/regulariser", {"type_resolution": "vol_presume", "motif": "x"}, stock)[0] == 403)
        c, b = api("GET", f"/inventaire/historique-comptages?date_debut={jour}&date_fin={jour}", jeton=resp)
        verif("C7 historique : tous les comptages, régularisation incluse",
              c == 200 and len(b["comptages"]) == int(sql("SELECT count(*) FROM comptages_stock"))
              and any(x["regularisation_type"] == "vol_presume" for x in b["comptages"]))
    try:
        import openpyxl  # type: ignore

        req = urllib.request.Request(BASE + f"/rapports/ecarts-comptage?date_debut={jour}&date_fin={jour}&format=xlsx",
                                     headers={"Authorization": "Bearer " + resp})
        classeur = openpyxl.load_workbook(io.BytesIO(urllib.request.urlopen(req).read()))
        lignes = list(classeur["écarts de comptage"].iter_rows(values_only=True))[1:]
        totaux = list(classeur["totaux"].iter_rows(values_only=True))[1]
        verif("C7 export : nb d'écarts et cumul = somme des lignes",
              totaux[0] == len(lignes) and totaux[1] == sum(r[7] for r in lignes), (totaux, len(lignes)))
    except ImportError:
        print("SKIP export xlsx (openpyxl absent)")

    # ------------------------------------------------------------------ concurrence (C10-02)
    for w in (1, 2, 4):
        sql("UPDATE stocks_sites SET quantite_stock = 1000 WHERE article_id=1 AND site_id=1")
        jetons = [resp, compta]

        def une(i, w=w):
            return vente(jetons[i % 2], {"site_id": 1, "mode_paiement": "especes", "lignes": [ligne(prix_unitaire=6500)]})

        with cf.ThreadPoolExecutor(w) as ex:
            res = list(ex.map(une, range(20)))
        codes = [c for c, _ in res]
        verif(f"C10-02 [DÉFAUT NOUVEAU] 20 ventes du même article, {w} client(s) simultané(s) : toutes acceptées",
              codes.count(201) == 20, {c: codes.count(c) for c in sorted(set(codes))})

    # ------------------------------------------------------------------ 5xx sur entrées aberrantes
    bruit = [
        ("vente 2,5 sur article entier", lambda: vente(compta, {"mode_paiement": "especes", "lignes": [ligne(quantite=2.5, prix_unitaire=6500)]})),
        ("vente 0,0001", lambda: vente(compta, {"mode_paiement": "especes", "lignes": [ligne(quantite=0.0001, prix_unitaire=6500)]})),
        ("vente prix 1e15", lambda: vente(compta, {"mode_paiement": "especes", "lignes": [ligne(prix_unitaire=1e15)]})),
        ("entrée de stock 1e12", lambda: api("POST", "/stock/entrees", {"article_id": 1, "quantite": 1e12}, stock)),
    ]
    for nom, appel in bruit:
        c, _ = appel()
        verif(f"ENTRÉE ABERRANTE [DÉFAUT NOUVEAU] {nom} : refus 4xx et non 500", c < 500, c)

    ok = sum(1 for _, o, _ in RESULTATS if o)
    print(f"\nBILAN : {ok} OK / {len(RESULTATS)} contrôles ; FAIL : {[n for n, o, _ in RESULTATS if not o]}")
    return 0 if ok == len(RESULTATS) else 1


if __name__ == "__main__":
    sys.exit(main())
